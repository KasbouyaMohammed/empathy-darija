"""Exploratory human validation (added after the frozen plan): one native Darija speaker, not involved in the
study, rates ~200 benchmark replies blind to system and policy. The sample targets the contested results:
  A. RQ2: Claude Sonnet 5.5 replies under P0, P1, P2 to Darija (Arabic-script) messages (complaints + info).
  B. RQ4: distilled LoRA (seed 0, no prompt) vs. Sonnet P2 vs. base Atlas-Chat-2B P2, Darija + Arabizi messages.
  C. RQ1: open models (Atlas-Chat-9B, Qwen3, Llama-3.1, Gemma-2) under P2 on Darija + Arabizi messages.
Only benchmark messages (our own texts) are shown; DarijaBanking human-written messages are not used.

    python scripts/18_human_eval.py build     # results/human_eval/human_eval_sheet.xlsx (+ hidden key)
    python scripts/18_human_eval.py score     # after the sheets come back: human_eval_filled.xlsx
                                              # (+ optional human_eval_filled_second_rater.xlsx) -> score_*.csv
"""
import sys
import warnings
from pathlib import Path

import numpy as np
import pandas as pd
from scipy.stats import spearmanr
from sklearn.metrics import cohen_kappa_score

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
OUT = ROOT / "results/human_eval"
SEED = 2026
N_RQ2_COMPLAINTS, N_RQ2_INFO, N_RQ4, N_RQ1_PER_SYSTEM = 20, 8, 25, 10

INSTRUCTIONS_AR = """تعليمات التقييم (عافاك قرا كلشي قبل ما تبدا)

كل سطر فيه رسالة ديال كليان كيكتب للشات بوت ديال بنكة مغربية، والجواب ديال الشات بوت.
ما كتعرفش شكون كتب الجواب (شي نموذج ذكاء اصطناعي). قيّم الجواب كيف كتشوفو نتا، بلا ما تحاول تخمن.
البنكة وهمية: [HOTLINE] و[N] هوما بلاصات فارغة عادية، ماشي غلطة.

عمّر 5 ديال الخانات الملونين لكل سطر:

1) Empathy (من 1 حتى 5): واش الجواب كيبين فهم واهتمام بالكليان؟
   1 = بارد ولا كيتجاهل الكليان، 3 = مؤدب ولكن عادي، 5 = فاهم الحالة ومهتم بصح.

2) Appropriateness (اختار وحدة): واش القدر ديال المشاعر فالجواب مناسب لحالة هاد الكليان؟
   too_little = الكليان معصب ولا خايف ولا مقلق والجواب ما دار ليه حتى اعتبار
   appropriate = القدر مناسب، وما كيعطلش الخطوات المستعجلة
   excessive = بزاف ديال الكلام العاطفي، ولا مكرر، ولا كيأخر الحل المستعجل، ولا سؤال عادي وجاوبو بالتعاطف بلا سبب

3) Helpfulness (من 1 حتى 5): واش الجواب كيعاون بصح؟
   1 = ما فيه حتى معاونة ولا معلومات غالطة، 3 = كيعاون شوية، 5 = كامل وواضح وتقدر تطبقو.

4) Same language (yes/no): واش الجواب مكتوب بنفس اللغة ونفس الكتابة ديال الكليان؟
   مثلا: الكليان كتب بالدارجة بالحروف اللاتينية (3, 7, 9) والجواب بالحروف العربية = no.
   الكليان كتب بالدارجة والجواب بالفصحى = no.

5) Naturalness (من 1 حتى 5، غير إلا كان الكليان كاتب بالدارجة): واش الجواب طبيعي بحال شي مغربي؟
   1 = ماشي طبيعي ولا ماشي دارجة، 5 = طبيعي مية فالمية. إلا الكليان ما كتبش بالدارجة خليها خاوية.

ملاحظات:
- قيّم كل سطر بوحدو. ما ترجعش تبدل أجوبة قدام باش "توازن".
- إلا شي جواب مقطوع ولا ما مفهومش، قيّمو كيف هو.
- تقدر تكتب ملاحظة فخانة comment (اختياري).
- الوقت التقريبي: 3 حتى 4 سوايع. تقدر تقسمها على بزاف ديال المرات.
شكرا بزاف على المساعدة!"""


def load_all():
    import glob
    fs = [ROOT / "data/processed/responses_claude.csv"] + [Path(f) for f in sorted(glob.glob(str(ROOT / "data/processed/responses_open_*.csv")))
                                                          if "pilot" not in f and "smoke" not in f]
    d = pd.concat([pd.read_csv(f) for f in fs], ignore_index=True)
    d = d[(d.set == "benchmark") & d.reply.notna() & (d.reply.astype(str).str.strip() != "")]
    if "excluded" in d:
        d = d[d.excluded.fillna(0) != 1]
    d["item_key"] = d.item_id + "|" + d.system + "|" + d.policy
    return d


def build():
    rng = np.random.default_rng(SEED)
    d = load_all()
    parts = []
    # A. RQ2: Sonnet, Darija (Arabic script), complaints (high/medium) and info requests, all three policies.
    a = d[(d.system == "claude-sonnet-5-5") & (d.variety == "ary")]
    comp = sorted(a[a.urgency.isin(["high", "medium"])].scenario_id.unique())
    info = sorted(a[a.urgency == "low"].scenario_id.unique())
    pick = list(rng.choice(comp, N_RQ2_COMPLAINTS, replace=False)) + list(rng.choice(info, N_RQ2_INFO, replace=False))
    parts.append(a[a.scenario_id.isin(pick) & a.policy.isin(["P0", "P1", "P2"])].assign(block="A_policies"))
    # B. RQ4: Darija + Arabizi, LoRA seed 0 vs Sonnet P2 vs base Atlas-2B P2, same messages.
    b = d[d.variety.isin(["ary", "ary_latn"])]
    arms = ((b.system == "lora2b_frac100_seed0") | ((b.system == "claude-sonnet-5-5") & (b.policy == "P2"))
            | ((b.system == "atlas2b") & (b.policy == "P2")))
    b = b[arms]
    full = b.groupby("item_id").system.nunique()
    items = sorted(full[full == 3].index)
    items = [i for i in items if i.split("_")[0] not in pick]           # no overlap with block A scenarios
    pick_b = list(rng.choice(items, N_RQ4, replace=False))
    parts.append(b[b.item_id.isin(pick_b)].assign(block="B_distillation"))
    # C. RQ1: open models under P2 on Darija + Arabizi messages not used above.
    c = d[d.system.isin(["atlas9b", "qwen3_8b", "llama31_8b", "gemma2_9b"]) & (d.policy == "P2")
          & d.variety.isin(["ary", "ary_latn"]) & ~d.scenario_id.isin(pick) & ~d.item_id.isin(pick_b)]
    for sy, g in c.groupby("system"):
        parts.append(g.loc[rng.choice(g.index, N_RQ1_PER_SYSTEM, replace=False)].assign(block="C_open_models"))
    s = pd.concat(parts, ignore_index=True)
    s = s.sample(frac=1, random_state=SEED).reset_index(drop=True)       # random order: blocks are mixed
    s["eval_id"] = [f"H{i:03d}" for i in range(1, len(s) + 1)]
    OUT.mkdir(parents=True, exist_ok=True)
    s[["eval_id", "item_key", "block", "system", "policy", "variety", "scenario_id", "category", "urgency"]] \
        .to_csv(OUT / "human_eval_key.csv", index=False)                  # hidden key: never give to the rater
    write_sheet(s, "human_eval_sheet.xlsx")
    write_sheet(s.head(50), "human_eval_sheet_second_rater_50.xlsx")   # optional overlap for inter-rater agreement
    print(f"{len(s)} items:", s.block.value_counts().to_dict(), "->", OUT / "human_eval_sheet.xlsx")


def write_sheet(s, name):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = Workbook()
    ins = wb.active
    ins.title = "Instructions"
    ins.sheet_view.rightToLeft = True
    for i, line in enumerate(INSTRUCTIONS_AR.split("\n"), 1):
        ins.cell(row=i, column=1, value=line).alignment = Alignment(wrap_text=True, horizontal="right")
    ins.column_dimensions["A"].width = 110
    ws = wb.create_sheet("Rate")
    head = ["eval_id", "customer_message", "chatbot_reply", "empathy_1_5", "appropriateness", "helpfulness_1_5",
            "same_language_yes_no", "naturalness_1_5_darija_only", "comment"]
    ws.append(head)
    for r in s.itertuples():
        ws.append([r.eval_id, r.message, r.reply, None, None, None, None, None, None])
    fill = PatternFill("solid", fgColor="FFF2CC")
    for col, width in zip("ABCDEFGHI", [9, 55, 75, 12, 18, 14, 16, 18, 25]):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        for c in row[:3]:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for c in row[3:8]:
            c.fill = fill
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.freeze_panes = "B2"
    n = len(s) + 1
    for rng_, formula in ((f"D2:D{n}", '"1,2,3,4,5"'), (f"F2:F{n}", '"1,2,3,4,5"'), (f"H2:H{n}", '"1,2,3,4,5"'),
                          (f"E2:E{n}", '"too_little,appropriate,excessive"'), (f"G2:G{n}", '"yes,no"')):
        dv = DataValidation(type="list", formula1=formula, allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(rng_)
    wb.save(OUT / name)


HUMAN = {"empathy_1_5": "empathy", "appropriateness": "appropriateness", "helpfulness_1_5": "helpfulness",
         "same_language_yes_no": "same_language", "naturalness_1_5_darija_only": "naturalness"}
APPROP = {"too_little": -1, "appropriate": 0, "excessive": 1}
SCALES = ["empathy", "helpfulness", "naturalness"]
N_BOOT = 2000


def read_filled(path):
    """One rater's sheet -> one row per eval_id; out-of-range or misspelled cells become NaN (and are counted)."""
    r = pd.read_excel(path, sheet_name="Rate").rename(columns=HUMAN)
    r["eval_id"] = r.eval_id.astype(str).str.strip()
    bad = {}
    for c in SCALES:
        v = pd.to_numeric(r[c], errors="coerce")
        ok = v.isin([1, 2, 3, 4, 5])
        bad[c] = int((r[c].notna() & ~ok).sum())
        r[c] = v.where(ok)
    a = r.appropriateness.astype(str).str.strip().str.lower().str.replace(" ", "_")
    r["approp"] = a.map(APPROP)
    bad["appropriateness"] = int((r.appropriateness.notna() & r.approp.isna()).sum())
    s = r.same_language.astype(str).str.strip().str.lower()
    r["same_language"] = s.map({"yes": 1, "no": 0})
    bad["same_language"] = int((s.ne("nan") & r.same_language.isna()).sum())
    r["appropriate"] = (r.approp == 0).astype(float).where(r.approp.notna())
    print(path.name, "rated:", r[SCALES + ["approp", "same_language"]].notna().sum().to_dict(), "invalid:", bad)
    return r[["eval_id", *SCALES, "approp", "appropriate", "same_language"]]


def boot_ci(f, *cols, seed=SEED):
    """Percentile 95% CI of statistic f over items resampled with replacement."""
    rng = np.random.default_rng(seed)
    n = len(cols[0])
    vals = []
    for _ in range(N_BOOT):
        i = rng.integers(0, n, n)
        vals.append(f(*(np.asarray(c)[i] for c in cols)))
    return np.nanpercentile(vals, [2.5, 97.5])


def paired(x, y, measure):
    """x - y on paired items: mean difference, bootstrap CI, Wilcoxon (scales) or exact McNemar (binary)."""
    from scipy.stats import binomtest, wilcoxon
    m = ~(np.isnan(x) | np.isnan(y))
    x, y = x[m], y[m]
    d = x - y
    if measure in ("appropriate", "same_language"):
        b, c = int(((x == 1) & (y == 0)).sum()), int(((x == 0) & (y == 1)).sum())
        p = binomtest(b, b + c).pvalue if b + c else 1.0
    else:
        p = wilcoxon(x, y).pvalue if np.any(d != 0) else 1.0
    lo, hi = boot_ci(np.mean, d) if len(d) else (np.nan, np.nan)
    return dict(n=len(d), mean_a=x.mean(), mean_b=y.mean(), diff=d.mean(), ci_lo=lo, ci_hi=hi, p=p)


def contrasts(h, block, by, pairs, groups=None):
    """Paired contrasts within a block; items are matched on `by` (scenario or message)."""
    from importlib import import_module
    holm = import_module("12_analysis").holm
    rows = []
    b = h[h.block == block]
    for gname, g in (groups(b) if groups else [("all", b)]):
        for a, z in pairs:
            wide = {arm: g[g.arm == arm].set_index(by) for arm in (a, z)}
            common = wide[a].index.intersection(wide[z].index)
            for m in SCALES + ["appropriate", "same_language"]:
                x, y = wide[a].loc[common, m].astype(float).values, wide[z].loc[common, m].astype(float).values
                if np.isnan(x).all() or np.isnan(y).all():
                    continue
                rows.append(dict(block=block, group=gname, a=a, b=z, measure=m, **paired(x, y, m)))
    t = pd.DataFrame(rows)
    t["p_holm"] = t.groupby("measure").p.transform(lambda p: holm(p.values))     # Holm per measure within the block
    return t


def score():
    """Exploratory human validation (not preregistered). Outputs results/human_eval/score_*.csv."""
    from importlib import import_module
    an = import_module("12_analysis")
    warnings.filterwarnings("ignore", category=RuntimeWarning)          # constant bootstrap resamples -> NaN, dropped
    warnings.filterwarnings("ignore", message=".*constant.*")
    key = pd.read_csv(OUT / "human_eval_key.csv")
    h = key.merge(read_filled(OUT / "human_eval_filled.xlsx"), on="eval_id", how="left")
    h["arm"] = np.where(h.system == "lora2b_frac100_seed0", "lora", h.system.str.replace("claude-sonnet-5-5", "sonnet") + "_" + h.policy)
    h["item_id"] = h.item_key.str.split("|").str[0]
    h["complaint"] = np.where(h.urgency == "low", "information", "complaint")

    # 1. Human vs. jury (and each judge) on the same replies.
    l = pd.read_csv(ROOT / "results/analysis/replies_long.csv")
    l = l[l.item_key.isin(h.item_key)]
    obj = l.groupby("item_key").fidelity.first()
    per_judge = {"jury": l.groupby("item_key")[["empathy", "helpfulness", "naturalness", "approp"]].median()}
    for j, g in l.groupby("judge"):
        per_judge[j] = g.set_index("item_key")[["empathy", "helpfulness", "naturalness", "approp"]]
    rows = []
    for j, s in per_judge.items():
        m = h.set_index("item_key").join(s, rsuffix="_j")
        for meas in SCALES:
            v = m[[meas, meas + "_j"]].dropna()
            if len(v) < 10:
                continue
            rho = lambda a, b: spearmanr(a, b).statistic
            lo, hi = boot_ci(rho, v[meas].values, v[meas + "_j"].values)
            rows.append(dict(judge=j, measure=meas, n=len(v), stat="spearman", value=rho(v[meas], v[meas + "_j"]), ci_lo=lo, ci_hi=hi))
        v = m[["appropriate", "approp_j"]].dropna()
        if len(v) < 10:
            continue
        ja = (v.approp_j == 0).astype(int)
        rows.append(dict(judge=j, measure="appropriate", n=len(v), stat="agreement", value=(v.appropriate == ja).mean()))
        rows.append(dict(judge=j, measure="appropriate", n=len(v), stat="cohen_kappa",
                         value=cohen_kappa_score(v.appropriate.astype(int), ja)))
    v = h.set_index("item_key").join(obj)[["same_language", "fidelity"]].dropna()
    rows.append(dict(judge="identifier", measure="same_language", n=len(v), stat="agreement", value=(v.same_language == v.fidelity).mean()))
    rows.append(dict(judge="identifier", measure="same_language", n=len(v), stat="cohen_kappa",
                     value=cohen_kappa_score(v.same_language.astype(int), v.fidelity.astype(int))))
    pd.DataFrame(rows).to_csv(OUT / "score_human_vs_judges.csv", index=False)
    jury_lab = np.sign(per_judge["jury"].approp).map({-1: "too_little", 0: "appropriate", 1: "excessive"})
    lab = h.set_index("item_key").approp.map({-1: "too_little", 0: "appropriate", 1: "excessive"})
    pd.crosstab(lab.rename("human"), jury_lab.reindex(lab.index).rename("jury")).to_csv(OUT / "score_approp_crosstab.csv")

    # 2. Descriptives per block, arm and message type.
    desc = h.groupby(["block", "arm", "complaint", "variety"]).agg(
        n=("eval_id", "size"), empathy=("empathy", "mean"), helpfulness=("helpfulness", "mean"),
        naturalness=("naturalness", "mean"), appropriate=("appropriate", "mean"),
        too_little=("approp", lambda x: (x == -1).mean()), excessive=("approp", lambda x: (x == 1).mean()),
        same_language=("same_language", "mean")).reset_index()
    desc.to_csv(OUT / "score_descriptives.csv", index=False)
    fid = h.set_index("item_key").join(obj).groupby(["block", "arm"]).agg(
        n=("eval_id", "size"), human_same_language=("same_language", "mean"), identifier_fidelity=("fidelity", "mean"),
        naturalness=("naturalness", "mean")).reset_index()
    fid.to_csv(OUT / "score_fidelity.csv", index=False)

    # 3. RQ2 (block A, paired by scenario) and RQ4 (block B, paired by message), Holm within each block.
    a = contrasts(h, "A_policies", "scenario_id", [("sonnet_P2", "sonnet_P1"), ("sonnet_P2", "sonnet_P0"), ("sonnet_P1", "sonnet_P0")],
                  groups=lambda b: b.groupby("complaint"))
    b = contrasts(h, "B_distillation", "item_id", [("lora", "sonnet_P2"), ("lora", "atlas2b_P2")])
    b["noninferior_0.3"] = np.where(b.b.eq("sonnet_P2") & b.measure.isin(SCALES), b.ci_lo > -0.3, np.nan)
    pd.concat([a, b]).to_csv(OUT / "score_contrasts.csv", index=False)

    # 4. Inter-rater agreement on the overlapping items (optional second rater).
    second = OUT / "human_eval_filled_second_rater.xlsx"
    if second.exists():
        import krippendorff
        r1 = h.set_index("eval_id")
        r2 = read_filled(second).set_index("eval_id")
        ids = r2.index.intersection(r1.index)
        rows = []
        for meas, level in (("empathy", "ordinal"), ("helpfulness", "ordinal"), ("naturalness", "ordinal"),
                            ("approp", "nominal"), ("same_language", "nominal")):
            x, y = r1.loc[ids, meas].astype(float), r2.loc[ids, meas].astype(float)
            m = x.notna() & y.notna()
            alpha = krippendorff.alpha(reliability_data=np.vstack([x, y]), level_of_measurement=level)
            k = (an.qwk(x[m] - 1, y[m] - 1, k=5) if level == "ordinal"
                 else cohen_kappa_score(x[m].astype(int), y[m].astype(int)))
            rows.append(dict(measure=meas, n=int(m.sum()), exact_agreement=(x[m] == y[m]).mean(),
                             kappa=k, kappa_type="quadratic" if level == "ordinal" else "cohen", krippendorff_alpha=alpha))
        pd.DataFrame(rows).to_csv(OUT / "score_interrater.csv", index=False)
        cols = ["empathy", "helpfulness", "naturalness", "approp", "same_language"]
        d = (r1.loc[ids, cols].astype(float) - r2.loc[ids, cols].astype(float)).abs()
        pd.DataFrame({"eval_id": ids, "identical": (d.fillna(0) == 0).all(axis=1).values,
                      "max_abs_diff_scales": d[SCALES].max(axis=1).values}).to_csv(OUT / "score_interrater_items.csv", index=False)
        print(pd.DataFrame(rows).round(3).to_string(index=False))
    print(pd.read_csv(OUT / "score_human_vs_judges.csv").round(3).to_string(index=False))
    print(pd.concat([a, b]).round(3).to_string(index=False))


if __name__ == "__main__":
    {"build": build, "score": score}[sys.argv[1]]()
